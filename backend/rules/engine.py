"""
规则引擎：使用 python-docx 读取 docx，调用 GB/T 9704 规则逐项检查。

格式提取要点（修复历史误报）：
  - 行距：固定值/最小值为 Length（EMU），须用 .pt 取磅值；倍数型保留倍数
  - 首行缩进：优先读 w:ind/@w:firstLineChars（字符单位，Word/WPS 界面的
    "首行缩进2字符"存的就是它），换算为磅值参与比较
  - 字体字号：沿 段落样式链 → 基于样式 → docDefaults 解析有效格式，
    不再只看第一个 run 的显式设置；中文字体名优先读 w:eastAsia
"""

from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Length
from docx.enum.text import WD_ALIGN_PARAGRAPH

from backend.core.models import Issue, IssueCategory, IssueSeverity
from backend.rules import gb_t9704 as gb


# ---------------------------------------------------------------------------
# 格式解析辅助（_rpr_font 在 gb_t9704 中定义，这里复用）
# ---------------------------------------------------------------------------

def _run_font(run) -> tuple[Optional[str], Optional[float]]:
    """run 的显式字体字号（读原始 XML，兼容 w:eastAsia）。"""
    return gb._rpr_font(run._r.find(qn("w:rPr")))


def _style_font(style) -> tuple[Optional[str], Optional[float]]:
    """样式自身 rPr 的字体字号。"""
    if style is None:
        return None, None
    return gb._rpr_font(style.element.find(qn("w:rPr")))


def _doc_default_font(doc) -> tuple[Optional[str], Optional[float]]:
    """文档默认字体字号（styles.xml 的 docDefaults/rPrDefault）。"""
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is None:
        return None, None
    rpr_default = dd.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return None, None
    return gb._rpr_font(rpr_default.find(qn("w:rPr")))


def _effective_para_font(para, doc) -> tuple[Optional[str], Optional[float]]:
    """解析段落的有效字体字号：
    run 显式设置 → 段落样式 → 样式链（basedOn）→ docDefaults。"""
    name = size = None

    # 1) 段落内 run 的显式设置
    for run in para.runs:
        rn, rs = _run_font(run)
        if name is None and rn:
            name = rn
        if size is None and rs:
            size = rs
        if name and size:
            return name, size

    # 2) 段落样式及样式链
    style = para.style
    while style is not None and (name is None or size is None):
        sn, ss = _style_font(style)
        if name is None and sn:
            name = sn
        if size is None and ss:
            size = ss
        style = style.base_style

    # 3) docDefaults
    if name is None or size is None:
        dn, ds = _doc_default_font(doc)
        if name is None and dn:
            name = dn
        if size is None and ds:
            size = ds

    return name, size


# ---------------------------------------------------------------------------
# 文档信息提取
# ---------------------------------------------------------------------------

def extract_document_info(doc: Document) -> gb.DocumentInfo:
    """从 python-docx Document 对象提取页面设置和段落信息"""
    section = doc.sections[0]  # 取第一个节

    # 页面尺寸（EMU → mm）
    page_w = gb._emu_to_mm(section.page_width) if section.page_width else gb.A4_WIDTH_MM
    page_h = gb._emu_to_mm(section.page_height) if section.page_height else gb.A4_HEIGHT_MM

    # 页边距（EMU → mm）
    margin_top = gb._emu_to_mm(section.top_margin) if section.top_margin else gb.MARGIN_TOP_MM
    margin_bottom = gb._emu_to_mm(section.bottom_margin) if section.bottom_margin else gb.MARGIN_BOTTOM_MM
    margin_left = gb._emu_to_mm(section.left_margin) if section.left_margin else gb.MARGIN_LEFT_MM
    margin_right = gb._emu_to_mm(section.right_margin) if section.right_margin else gb.MARGIN_RIGHT_MM

    # 提取段落信息
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text

        # 有效字体字号（含样式继承解析）
        font_name, font_size = _effective_para_font(para, doc)

        # 段落对齐
        alignment = None
        if para.alignment is not None:
            alignment = int(para.alignment)

        pf = para.paragraph_format

        # 行距：固定值/最小值为 Length → 磅值；倍数型 → 倍数
        line_spacing = None
        if pf.line_spacing is not None:
            if isinstance(pf.line_spacing, Length):
                line_spacing = pf.line_spacing.pt
            else:
                line_spacing = float(pf.line_spacing)

        # 首行缩进：优先字符单位 firstLineChars（百倍字符数），
        # 换算为磅值；否则用磅值
        first_line_indent = None
        ppr = para._p.pPr
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        if ind is not None and ind.get(qn("w:firstLineChars")):
            chars = int(ind.get(qn("w:firstLineChars"))) / 100.0
            # 字符数 × 段落字号（缺省按三号 16pt）
            first_line_indent = chars * (font_size or gb.SIZE_SAN_HAO)
        elif pf.first_line_indent is not None:
            first_line_indent = pf.first_line_indent.pt

        paragraphs.append({
            "text": text,
            "font_name": font_name,
            "font_size": font_size,
            "alignment": alignment,
            "line_spacing": line_spacing,
            "first_line_indent": first_line_indent,
        })

    return gb.DocumentInfo(
        page_width_mm=page_w,
        page_height_mm=page_h,
        margin_top_mm=margin_top,
        margin_bottom_mm=margin_bottom,
        margin_left_mm=margin_left,
        margin_right_mm=margin_right,
        paragraphs=paragraphs,
    )


# ---------------------------------------------------------------------------
# 主引擎
# ---------------------------------------------------------------------------

class RuleEngine:
    """GB/T 9704 格式规则检查引擎"""

    def __init__(self, strictness: str = "normal"):
        self.strictness = strictness

    def check_document(self, doc_path: str | Path) -> list[Issue]:
        """检查整个文档，返回所有问题列表"""
        doc_path = Path(doc_path)
        if not doc_path.exists():
            raise FileNotFoundError(f"文件不存在: {doc_path}")

        doc = Document(str(doc_path))
        doc_info = extract_document_info(doc)

        all_issues: list[Issue] = []

        # 1. 页面幅面检查
        all_issues.extend(gb.check_page_size(doc_info, self.strictness))

        # 2. 页边距检查
        all_issues.extend(gb.check_page_margins(doc_info, self.strictness))

        # 3. 段落格式检查
        total = len(doc_info.paragraphs)
        for idx, para_info in enumerate(doc_info.paragraphs):
            text = para_info.get("text", "").strip()
            if not text:
                continue

            role = gb._identify_paragraph_role(text, para_info, idx, total)

            # 段落基础格式检查
            all_issues.extend(
                gb.check_paragraph_format(para_info, idx, role, self.strictness)
            )

            # 首行缩进检查（仅正文）
            if role == "main_body":
                indent_issue = gb.check_first_line_indent(
                    para_info, idx, text, self.strictness
                )
                if indent_issue:
                    all_issues.append(indent_issue)

        # 4. 页码检查（读 footer XML）
        all_issues.extend(gb.check_page_numbers(doc_path, self.strictness))

        return all_issues

    def check_document_object(self, doc: Document) -> list[Issue]:
        """从已打开的 Document 对象检查（避免重复打开；不含页码检查）"""
        doc_info = extract_document_info(doc)

        all_issues: list[Issue] = []

        all_issues.extend(gb.check_page_size(doc_info, self.strictness))
        all_issues.extend(gb.check_page_margins(doc_info, self.strictness))

        total = len(doc_info.paragraphs)
        for idx, para_info in enumerate(doc_info.paragraphs):
            text = para_info.get("text", "").strip()
            if not text:
                continue

            role = gb._identify_paragraph_role(text, para_info, idx, total)
            all_issues.extend(
                gb.check_paragraph_format(para_info, idx, role, self.strictness)
            )
            if role == "main_body":
                indent_issue = gb.check_first_line_indent(
                    para_info, idx, text, self.strictness
                )
                if indent_issue:
                    all_issues.append(indent_issue)

        return all_issues


# ---------------------------------------------------------------------------
# 便捷函数
# ---------------------------------------------------------------------------

def check_document_rules(doc_path: str | Path, strictness: str = "normal") -> list[Issue]:
    """便捷函数：检查文档格式"""
    engine = RuleEngine(strictness=strictness)
    return engine.check_document(doc_path)
